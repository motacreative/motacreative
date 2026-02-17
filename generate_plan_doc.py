#!/usr/bin/env python3
"""Generate a Word document from the migration plan."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# -- Styles --
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# Title
title = doc.add_heading('SureCart to FluentCart Migration Plugin', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('')

# --- Context ---
doc.add_heading('Context', level=1)
doc.add_paragraph(
    'The user needs to migrate their eCommerce store from SureCart to FluentCart. '
    'SureCart stores data in a cloud-hosted API (not locally in WordPress), while '
    'FluentCart uses local custom database tables. No existing migration tool exists '
    'between these two platforms. This plugin will run on a WordPress site where both '
    'SureCart and FluentCart are installed, reading from SureCart\'s PHP models and '
    'writing directly to FluentCart\'s database via its internal models.'
)

# --- Plugin Structure ---
doc.add_heading('Plugin Structure', level=1)
structure = """surecart-to-fluentcart/
├── surecart-to-fluentcart.php              # Main bootstrap, dependency checks, autoloader
├── includes/
│   ├── class-migration-activator.php       # Creates mapping table on activation
│   ├── class-admin-page.php                # Admin menu page + UI rendering
│   ├── class-ajax-handler.php              # AJAX batch dispatcher
│   ├── class-migration-logger.php          # Error/progress logging
│   ├── class-id-mapper.php                 # SureCart ID <-> FluentCart ID mapping
│   └── migrators/
│       ├── class-base-migrator.php         # Abstract base with batch/pagination logic
│       ├── class-product-migrator.php      # Products + variants + images
│       ├── class-customer-migrator.php     # Customers + addresses
│       └── class-order-migrator.php        # Orders + items + transactions
└── assets/
    ├── js/migration-admin.js               # AJAX loop, progress bars
    └── css/migration-admin.css             # Admin page styling"""
p = doc.add_paragraph()
run = p.add_run(structure)
run.font.name = 'Consolas'
run.font.size = Pt(9)

# --- How It Works ---
doc.add_heading('How It Works', level=1)
steps = [
    'Plugin adds a page under Tools > SureCart to FluentCart',
    'Admin UI shows dependency status, migration controls, and progress bars',
    'User clicks "Start Migration" — JS sends sequential AJAX requests in batches',
    'Migration runs in order: Products → Customers → Orders (orders depend on the other two)',
    'Each batch processes N records (configurable, default 20), skipping already-migrated items',
    'A mapping table (sc_fct_migration_map) tracks SureCart ID → FluentCart ID for cross-referencing',
    'Errors are logged per-record without stopping the batch',
]
for i, step in enumerate(steps, 1):
    doc.add_paragraph(f'{i}. {step}')

# --- Data Source ---
doc.add_heading('Data Source: SureCart PHP Models', level=1)
doc.add_paragraph(
    'SureCart data lives in their cloud API. We read it using their PHP models (SureCart\\Models\\*):'
)
models = [
    "Product::with(['prices', 'product_medias', 'product_medias.media', 'product_collections'])->paginate([...])",
    "Customer::with(['billing_address', 'shipping_address'])->paginate([...])",
    "Order::with(['checkout', 'checkout.customer', 'checkout.line_items', 'checkout.line_items.price', 'checkout.charges'])->paginate([...])",
]
for m in models:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(m)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)

doc.add_paragraph('Note: Only one-time purchases — no subscription/recurring price handling needed.')
doc.add_paragraph('Each call is a synchronous HTTP request — safe in admin AJAX context.')

# --- Data Destination ---
doc.add_heading('Data Destination: FluentCart Internal Models', level=1)
doc.add_paragraph(
    'Write directly to FluentCart\'s DB using wp_insert_post() for products (CPT fluent-products) '
    'and Eloquent-style model queries for custom tables:'
)
tables = [
    'fct_product_details / fct_product_variations — product pricing and variants',
    'fct_customers / fct_customer_addresses — customer profiles',
    'fct_orders / fct_order_items / fct_order_transactions / fct_order_addresses — orders',
]
for t in tables:
    doc.add_paragraph(t, style='List Bullet')
doc.add_paragraph('All monetary values are BIGINT in cents in both systems (no conversion needed).')

# --- Mapping Table ---
doc.add_heading('Mapping Table', level=1)
sql = """CREATE TABLE {prefix}sc_fct_migration_map (
  id BIGINT UNSIGNED AUTO_INCREMENT PRIMARY KEY,
  entity_type VARCHAR(50) NOT NULL,      -- 'product', 'customer', 'order', 'variant', 'collection'
  surecart_id VARCHAR(255) NOT NULL,     -- SureCart UUID strings
  fluentcart_id BIGINT UNSIGNED NOT NULL,
  meta TEXT NULL,                         -- Optional JSON
  created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
  UNIQUE KEY (entity_type, surecart_id)
);"""
p = doc.add_paragraph()
run = p.add_run(sql)
run.font.name = 'Consolas'
run.font.size = Pt(9)

# --- Field Mappings: Products ---
doc.add_heading('Field Mappings', level=1)
doc.add_heading('Products (SureCart → FluentCart)', level=2)

product_rows = [
    ('SureCart', 'FluentCart', 'Notes'),
    ('name', 'wp_posts.post_title', 'Direct'),
    ('description', 'wp_posts.post_content', 'HTML'),
    ('slug', 'wp_posts.post_name', 'URL slug'),
    ('product_medias.media', 'WordPress media library', 'Sideload via media_handle_sideload(), set featured + gallery'),
    ('product_collections', 'product-categories taxonomy', 'Create terms, assign to post'),
    ('Each price object', 'fct_product_variations row', 'price.amount → item_price, payment_type always onetime'),
    ('Min/max of prices', 'fct_product_details.min_price/max_price', 'Computed'),
    ('archived flag', 'post_status', 'false → publish, true → draft'),
]

table = doc.add_table(rows=len(product_rows), cols=3)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, row_data in enumerate(product_rows):
    for j, cell_text in enumerate(row_data):
        cell = table.cell(i, j)
        cell.text = cell_text
        if i == 0:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

# --- Field Mappings: Customers ---
doc.add_heading('Customers (SureCart → FluentCart)', level=2)

customer_rows = [
    ('SureCart', 'FluentCart fct_customers', 'Notes'),
    ('email', 'email', 'Dedup check first'),
    ('first_name / last_name', 'first_name / last_name', 'Direct'),
    ('billing_address.*', 'country, city, state, postcode', 'Top-level fields'),
    ('WordPress user lookup', 'user_id', 'get_user_by(\'email\', ...)'),
    ('Both addresses', 'fct_customer_addresses rows', 'billing + shipping types'),
]

table = doc.add_table(rows=len(customer_rows), cols=3)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, row_data in enumerate(customer_rows):
    for j, cell_text in enumerate(row_data):
        cell = table.cell(i, j)
        cell.text = cell_text
        if i == 0:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

# --- Field Mappings: Orders ---
doc.add_heading('Orders (SureCart → FluentCart)', level=2)

order_rows = [
    ('SureCart', 'FluentCart fct_orders', 'Notes'),
    ('checkout.customer', 'customer_id', 'Resolved via mapping table'),
    ('checkout.total_amount', 'total_amount', 'Cents, direct'),
    ('checkout.tax_amount', 'tax_total', 'Cents'),
    ('checkout.discount_amount', 'coupon_discount_total', 'Cents'),
    ('checkout.status', 'status + payment_status', 'Mapped (see below)'),
    ('checkout.currency', 'currency', 'Uppercase 3-letter'),
    ('Each line_item', 'fct_order_items row', 'Product/variant resolved via mapping'),
    ('Each charge', 'fct_order_transactions row', 'Amount, status, payment mode'),
    ('checkout.billing_address', 'fct_order_addresses (billing)', ''),
    ('checkout.shipping_address', 'fct_order_addresses (shipping)', ''),
]

table = doc.add_table(rows=len(order_rows), cols=3)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, row_data in enumerate(order_rows):
    for j, cell_text in enumerate(row_data):
        cell = table.cell(i, j)
        cell.text = cell_text
        if i == 0:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

# --- Status Mapping ---
doc.add_heading('Status Mapping', level=2)

status_rows = [
    ('SureCart Status', 'FluentCart Status', 'Payment Status'),
    ('paid', 'completed', 'paid'),
    ('pending', 'on-hold', 'pending'),
    ('payment_failed', 'failed', 'failed'),
    ('canceled', 'canceled', 'failed'),
    ('refunded', 'canceled', 'refunded'),
    ('partially_refunded', 'completed', 'partially_refunded'),
]

table = doc.add_table(rows=len(status_rows), cols=3)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, row_data in enumerate(status_rows):
    for j, cell_text in enumerate(row_data):
        cell = table.cell(i, j)
        cell.text = cell_text
        if i == 0:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

# --- Implementation Steps ---
doc.add_heading('Implementation Steps', level=1)

impl_steps = [
    ('Plugin scaffold', 'main file with constants, dependency checks (SureCart + FluentCart active), autoloader, activation hook'),
    ('Mapping infrastructure', 'MigrationActivator (dbDelta), IdMapper (CRUD), MigrationLogger'),
    ('Base migrator', 'abstract class with fetchBatch() / migrateOne() / processBatch() pattern, idempotency via IdMapper'),
    ('Product migrator', 'wp_insert_post() + ProductDetail + ProductVariation rows + image sideloading + taxonomy mapping'),
    ('Customer migrator', 'dedup by email, create Customer + CustomerAddresses rows, link WordPress user'),
    ('Order migrator', 'resolve customer/product/variant IDs from mapper, create Order + OrderItem + OrderTransaction + OrderAddress rows'),
    ('Admin UI', 'PHP page under Tools menu, JS AJAX loop with progress bars, CSS styling'),
    ('AJAX handler', 'nonce verification, capability checks (manage_options), dispatch to correct migrator, return progress JSON'),
]

for i, (title, desc) in enumerate(impl_steps, 1):
    p = doc.add_paragraph()
    run = p.add_run(f'{i}. {title}')
    run.bold = True
    p.add_run(f' — {desc}')

# --- Key Design Decisions ---
doc.add_heading('Key Design Decisions', level=1)

decisions = [
    ('Direct DB writes over REST API', 'faster, no auth overhead, matches FluentCart\'s own migration patterns'),
    ('AJAX batches over WP-Cron', 'gives real-time progress feedback, simpler to implement'),
    ('Idempotent', 're-running skips already-mapped records, safe to resume after failure'),
    ('Lock mechanism', 'transient-based lock prevents concurrent runs'),
    ('Image sideloading', 'downloads SureCart CDN images into local media library (slow but necessary)'),
]

for title, desc in decisions:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(title)
    run.bold = True
    p.add_run(f' — {desc}')

# --- Verification ---
doc.add_heading('Verification', level=1)

verification = [
    'Activate plugin on a WordPress site with both SureCart and FluentCart active',
    'Navigate to Tools > SureCart to FluentCart',
    'Verify dependency indicators show green for both plugins',
    'Run migration with a small batch size (5) first',
    'Check FluentCart admin: products appear with correct names, prices, images, and categories',
    'Check FluentCart admin: customers appear with correct emails, names, and addresses',
    'Check FluentCart admin: orders appear linked to correct customers and products, with correct totals and statuses',
    'Run migration again — verify no duplicates are created (idempotency)',
    'Test "Reset" functionality clears mapping data',
]

for i, step in enumerate(verification, 1):
    doc.add_paragraph(f'{i}. {step}')

# Save
output_path = '/home/user/motacreative/SureCart-to-FluentCart-Migration-Plan.docx'
doc.save(output_path)
print(f'Document saved to: {output_path}')
